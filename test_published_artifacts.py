"""Opt-in smoke tests of the exact published artifacts; never train or refresh."""

import csv
import hashlib
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook


@unittest.skipUnless(os.environ.get("DID_RUN_ARTIFACT_SMOKE") == "1", "Set DID_RUN_ARTIFACT_SMOKE=1 for real artifact smoke tests")
class PublishedArtifactTests(unittest.TestCase):
    def test_all_four_original_bytes_match_published_provenance(self):
        root = Path(__file__).resolve().parent / "artifacts"
        provenance = json.loads((root / "provenance.json").read_text())
        self.assertEqual(len(provenance["artifacts"]), 4)
        for name, expected in provenance["artifacts"].items():
            path = root / name
            self.assertEqual(path.stat().st_size, expected["size"])
            with path.open("rb") as content:
                self.assertEqual(hashlib.file_digest(content, "sha256").hexdigest(), expected["sha256"])

    def test_real_saved_model_inference_without_neo4j(self):
        import effort_prediction as prediction
        artifact = prediction._load_model_artifact(
            str(prediction.DEFAULT_MODEL_PATH), prediction.DEFAULT_MODEL_PATH.stat().st_mtime_ns
        )
        with patch("effort_prediction._read_query", side_effect=AssertionError("Unexpected Neo4j access")):
            result = prediction.predict_record(artifact["history"][-1], prediction.DEFAULT_MODEL_PATH, as_of_date="2026-10-10")
        self.assertEqual(result["model_version"], "did-effort-ridge-v3-robust")
        self.assertGreaterEqual(result["p50_hours"], 0)
        self.assertLessEqual(result["p50_hours"], result["p80_hours"])
        self.assertLessEqual(result["p80_hours"], result["p90_hours"])
        self.assertEqual(len(artifact["history"]), 32705)

    def test_real_snapshot_api_uploads_and_excel_without_llm_or_refresh(self):
        import tlf_person_allocation as allocation
        from app import app
        history, _, snapshot = allocation.load_person_tlf_snapshot()
        self.assertEqual(len(history), 1049456)
        row = next(row for row in history if row.get("tlf_name") and row.get("team_lead_name")
                   and any(allocation._assignment_matches(row.get(role), row.get("person"))
                           for role in ("generation", "qc")))
        lead = row["team_lead_name"].strip()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Title", "Type", "Source Datasets"])
        writer.writerow([row["tlf_name"], row.get("tlf_type", ""), row.get("tlf_source", "")])
        tlf_csv = output.getvalue().encode("utf-8")
        data_csv = b"SDTM/ADaM,Domain/Dataset Name\nADaM,ADAE\n"
        with tempfile.TemporaryDirectory() as cache_directory, \
             patch("effort_prediction._prediction_artifact_cache_dir", return_value=Path(cache_directory)), \
             patch("vox_client.chat", side_effect=AssertionError("Unexpected LLM call")), \
             patch("du_team_recommendation._read_query", side_effect=AssertionError("Unexpected Neo4j query")), \
             patch("tlf_person_allocation._read_query", side_effect=AssertionError("Unexpected Neo4j query")):
            client = TestClient(app)
            options = client.get("/api/planning/recommend/options")
            self.assertEqual(set(options.json()["ta_leads"]), {"Shen, Henry", "Iyer, Priya Venkatesan"})
            self.assertEqual(options.json()["snapshot"]["version"], "du-team-history-v4-ta-group-hierarchy")
            recommended = client.post("/api/planning/recommend", data={"ta_lead_name": "Shen, Henry"},
                                      files={"primary_file": ("TLF.csv", tlf_csv)})
            self.assertEqual(recommended.status_code, 200, recommended.text[:200])
            self.assertGreater(recommended.json()["candidate_team_count"], 0)
            self.assertEqual(len(recommended.json()["recommendations"]), 3)
            allocated = client.post("/api/planning/allocate", data={"team_lead_name": lead},
                                    files={"primary_file": ("TLF.csv", tlf_csv), "data_csv_file": ("Data.csv", data_csv)})
            self.assertEqual(allocated.status_code, 200, allocated.text[:200])
            result = allocated.json()
            self.assertEqual(result["team_match"]["team_lead_name"], lead)
            self.assertEqual(result["snapshot"]["version"], snapshot["version"])
            self.assertEqual(len(result["allocations"]), 1)
            item = result["allocations"][0]
            self.assertTrue(item["generation"]["primary"] or item["qc"]["primary"])
            if item["generation"]["primary"] and item["qc"]["primary"]:
                distinct = item["generation"]["primary"]["person"] != item["qc"]["primary"]["person"]
                self.assertTrue(distinct or item["lead_review_required"])
            excel = client.post("/api/planning/allocate/excel", json={"result": result})
            self.assertEqual(excel.status_code, 200)
            workbook = load_workbook(io.BytesIO(excel.content), read_only=True)
            self.assertGreater(workbook.active.max_row, 1)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
