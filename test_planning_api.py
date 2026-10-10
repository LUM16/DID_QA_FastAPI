"""Planning API contract, validation and Q&A coexistence regressions."""

import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import effort_prediction as prediction
from app import app


class PlanningAPITests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        import planning_api
        planning_api.JOBS.clear()

    @patch("planning_api.load_history_snapshot")
    def test_recommendation_options_are_snapshot_restricted(self, load):
        load.return_value = ([
            {"organization_mappings": [{"ta_lead": name, "group_lead": "G"}]}
            for name in ["Other", "Shen, Henry", "Iyer, Priya Venkatesan"]
        ], [], {"version": "v4"})
        response = self.client.get("/api/planning/recommend/options")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(set(response.json()["ta_leads"]), {"Shen, Henry", "Iyer, Priya Venkatesan"})

    @patch("planning_api.load_person_tlf_snapshot")
    def test_allocation_options_include_all_exact_team_leads(self, load):
        load.return_value = ([{"team_lead_name": "Z, Other"}, {"team_lead_name": "A, Example"}], [], {})
        response = self.client.get("/api/planning/allocate/options")
        self.assertEqual(response.json()["team_leads"], ["A, Example", "Z, Other"])

    @patch("planning_api.recommend_uploaded_scope")
    def test_csv_upload_adapts_to_published_business_api(self, recommend):
        recommend.return_value = {"candidate_team_count": 1, "recommendations": []}
        response = self.client.post("/api/planning/recommend", data={"ta_lead_name": "Shen, Henry"},
                                    files={"primary_file": ("TLF.csv", b"Name,Type\nTable,Table\n", "text/csv")})
        self.assertEqual(response.status_code, 200)
        name, primary, data = recommend.call_args.args
        self.assertEqual(name, "Shen, Henry")
        self.assertEqual(primary.name, "TLF.csv")
        self.assertIn(b"Table", primary.getvalue())
        self.assertIsNone(data)

    def test_recommendation_workflows_accept_tlf_only_csv_and_workbook_without_llm(self):
        book = Workbook()
        book.active.title = "TLF"
        book.active.append(["Title", "Type", "Source Datasets"])
        book.active.append(["AE Summary", "Table", "ADAE"])
        workbook = io.BytesIO()
        book.save(workbook)
        book.close()
        history = [{"organization_mappings": [{"ta_lead": "Shen, Henry", "group_lead": "G"}]}]
        with patch("du_team_recommendation.load_history_snapshot", return_value=(history, [], {})), \
             patch("du_team_recommendation.recommend_teams", return_value={"recommendations": []}) as rank, \
             patch("vox_client.chat", side_effect=AssertionError("Unexpected LLM call")), \
             patch("du_team_recommendation._read_query", side_effect=AssertionError("Unexpected Neo4j query")):
            for route in ["/api/planning/recommend", "/api/planning/jobs/recommend"]:
                for filename, content in [
                    ("TLF.csv", b"Title,Type,Source Datasets\nAE Summary,Table,ADAE\n"),
                    ("scope.xlsx", workbook.getvalue()),
                ]:
                    with self.subTest(route=route, filename=filename):
                        response = self.client.post(route, data={"ta_lead_name": "Shen, Henry"},
                                                    files={"primary_file": (filename, content)})
                        self.assertEqual(response.status_code, 202 if "/jobs/" in route else 200)
                        if "/jobs/" in route:
                            self.assertEqual(self.client.get(
                                f"/api/planning/jobs/{response.json()['job_id']}"
                            ).json()["status"], "complete")
                        self.assertEqual(rank.call_args.args[0]["tlfs"][0]["name"], "AE Summary")
                        self.assertEqual(rank.call_args.args[0]["sdtms"], [])

    def test_recommendation_contract_has_no_data_upload_field(self):
        schema = app.openapi()
        for route in ["/api/planning/recommend", "/api/planning/jobs/recommend"]:
            body = schema["paths"][route]["post"]["requestBody"]["content"]["multipart/form-data"]["schema"]
            properties = schema["components"]["schemas"][body["$ref"].split("/")[-1]]["properties"]
            self.assertEqual(set(properties), {"ta_lead_name", "primary_file"})

    def test_single_csv_and_tlf_only_workbook_allocation_and_excel_without_llm_or_database(self):
        book = Workbook()
        book.active.title = "TLF"
        book.active.append(["Title", "Type", "Source Datasets"])
        book.active.append(["AE Summary", "T", "ADAE"])
        workbook = io.BytesIO()
        book.save(workbook)
        book.close()
        with patch("tlf_person_allocation.load_person_tlf_snapshot",
                   return_value=([{"team_lead_name": "Exact, Lead"}], [], {})), \
             patch("tlf_person_allocation.allocate_tlf_people", return_value={
                 "allocations": [{
                     "tlf": {"name": "AE Summary", "type": "T", "source": "ADAE"}, "group": "Source: ADAE",
                     "generation": {"primary": None, "backups": []},
                     "qc": {"primary": None, "backups": []},
                     "lead_review_required": True, "fallback_reason": "Review",
                 }],
             }) as rank, \
             patch("vox_client.chat", side_effect=AssertionError("Unexpected LLM call")) as llm, \
             patch("tlf_person_allocation._read_query", side_effect=AssertionError("Unexpected Neo4j query")) as query:
            for route in ["/api/planning/allocate", "/api/planning/jobs/allocate"]:
                for filename, content in [
                    ("TLF.csv", b"Title,Type,Source Datasets\nAE Summary,T,ADAE\n"),
                    ("TLF-only.xlsx", workbook.getvalue()),
                ]:
                    with self.subTest(route=route, filename=filename):
                        response = self.client.post(route, data={"team_lead_name": "Exact, Lead"},
                                                    files={"primary_file": (filename, content)})
                        self.assertEqual(response.status_code, 202 if "/jobs/" in route else 200)
                        if "/jobs/" in route:
                            response = self.client.get(f"/api/planning/jobs/{response.json()['job_id']}")
                            self.assertEqual(response.json()["status"], "complete")
                            result = response.json()["result"]
                        else:
                            result = response.json()
                        self.assertEqual(rank.call_args.args[0], [{"name": "AE Summary", "type": "T", "source": "ADAE"}])
                        self.assertEqual(rank.call_args.kwargs["group_keys"], ["Source: ADAE"])
                        exported = self.client.post("/api/planning/allocate/excel", json={"result": result})
                        self.assertEqual(exported.status_code, 200)
                        exported_book = load_workbook(io.BytesIO(exported.content), read_only=True)
                        self.assertEqual(exported_book.active.max_row, 2)
                        exported_book.close()
            llm.assert_not_called()
            query.assert_not_called()

    def test_allocation_contract_has_no_secondary_data_upload(self):
        schema = app.openapi()
        for route in ["/api/planning/allocate", "/api/planning/jobs/allocate"]:
            body = schema["paths"][route]["post"]["requestBody"]["content"]["multipart/form-data"]["schema"]
            form_schema = schema["components"]["schemas"][body["$ref"].split("/")[-1]]
            self.assertEqual(set(form_schema["properties"]), {"team_lead_name", "primary_file"})
            self.assertEqual(set(form_schema["required"]), {"team_lead_name", "primary_file"})
        self.assertEqual(set(schema["components"]["schemas"]["AllocationExport"]["properties"]), {"result"})

    @patch("planning_api.allocate_uploaded_tlfs")
    def test_xlsx_and_single_csv_allocation_routes(self, allocate):
        allocate.return_value = {"allocations": []}
        for files in [
            {"primary_file": ("scope.xlsx", b"mocked workbook", "application/octet-stream")},
            {"primary_file": ("TLF.csv", b"Name\nA\n")},
        ]:
            response = self.client.post("/api/planning/allocate", data={"team_lead_name": "Exact, Lead"}, files=files)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(allocate.call_args.args[0], "Exact, Lead")
        self.assertEqual(len(allocate.call_args.args), 2)

    @patch("planning_api.allocate_uploaded_tlfs")
    def test_invalid_input_is_clear_422(self, allocate):
        allocate.side_effect = ValueError("Select an exact Team_Lead_Name.")
        response = self.client.post("/api/planning/allocate", data={"team_lead_name": "Partial"},
                                    files={"primary_file": ("TLF.csv", b"Name\nA\n")})
        self.assertEqual(response.status_code, 422)
        self.assertIn("exact", response.json()["detail"])

    @patch("planning_api.load_history_snapshot", side_effect=FileNotFoundError("missing"))
    def test_missing_snapshot_is_service_error_not_live_refresh(self, load):
        self.assertEqual(self.client.get("/api/planning/recommend/options").status_code, 503)

    def test_bad_extension_empty_and_missing_upload(self):
        for filename, content in [("bad.txt", b"x"), ("empty.csv", b"")]:
            response = self.client.post("/api/planning/recommend", data={"ta_lead_name": "Shen, Henry"},
                                        files={"primary_file": (filename, content)})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.post("/api/planning/allocate").status_code, 422)

    def test_size_limit(self):
        with patch("planning_api.MAX_UPLOAD_BYTES", 2):
            response = self.client.post("/api/planning/recommend", data={"ta_lead_name": "Shen, Henry"},
                                        files={"primary_file": ("TLF.csv", b"too big")})
        self.assertEqual(response.status_code, 413)

    @patch("planning_api.recommend_uploaded_scope")
    def test_background_job_results_and_progress_contract(self, recommend):
        recommend.return_value = {"candidate_team_count": 3, "recommendations": []}
        created = self.client.post("/api/planning/jobs/recommend", data={"ta_lead_name": "Shen, Henry"},
                                   files={"primary_file": ("TLF.csv", b"Title\nA\n")})
        self.assertEqual(created.status_code, 202)
        polled = self.client.get(f"/api/planning/jobs/{created.json()['job_id']}")
        self.assertEqual(polled.json()["status"], "complete")
        self.assertEqual(polled.json()["result"]["candidate_team_count"], 3)
        self.assertGreaterEqual(polled.json()["elapsed_seconds"], 0)

    @patch("planning_api.allocate_uploaded_tlfs", side_effect=ValueError("Select an exact name"))
    def test_background_job_validation_errors_and_expiry(self, allocate):
        created = self.client.post("/api/planning/jobs/allocate", data={"team_lead_name": "Partial"},
                                   files={"primary_file": ("TLF.csv", b"Title\nA\n")})
        polled = self.client.get(f"/api/planning/jobs/{created.json()['job_id']}")
        self.assertEqual(polled.json()["status"], "error")
        self.assertEqual(polled.json()["error_status"], 422)
        self.assertEqual(self.client.get("/api/planning/jobs/not-found").status_code, 404)
        with patch("planning_api.JOB_TTL_SECONDS", -1):
            self.assertEqual(self.client.get(f"/api/planning/jobs/{created.json()['job_id']}").status_code, 404)

    def test_job_queue_is_bounded(self):
        from fastapi import BackgroundTasks
        import planning_api
        with patch("planning_api.MAX_JOBS", 1):
            planning_api._new_job(BackgroundTasks(), lambda: {})
            response = self.client.post("/api/planning/jobs/recommend", data={"ta_lead_name": "Shen, Henry"},
                                        files={"primary_file": ("TLF.csv", b"Title\nA\n")})
        self.assertEqual(response.status_code, 429)

    def test_example_and_excel_export(self):
        example = self.client.get("/api/planning/example")
        self.assertEqual(example.status_code, 200)
        book = load_workbook(io.BytesIO(example.content), read_only=True)
        self.assertTrue({"TLF", "Data"}.issubset(set(book.sheetnames)))
        book.close()
        result = {
            "team_lead_name": "Exact, Lead",
            "allocations": [{
                "tlf": {"name": "A", "type": "Table", "source": "ADAE"}, "group": "AE",
                "generation": {"primary": None, "backups": []},
                "qc": {"primary": None, "backups": []},
                "lead_review_required": True, "fallback_reason": "Review",
            }],
        }
        response = self.client.post("/api/planning/allocate/excel", json={"result": result})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertGreater(book.active.max_row, 1)
        book.close()
        self.assertEqual(self.client.post("/api/planning/allocate/excel", json={"result": {}}).status_code, 422)

    def test_qa_and_visualization_routes_remain_present(self):
        paths = set(app.openapi()["paths"])
        self.assertTrue({"/api/query", "/api/feedback", "/api/export"}.issubset(paths))
        page = self.client.get("/")
        self.assertIn('id="chat"', page.text)
        self.assertIn("Allocate TLF People", page.text)

    @patch("app.agent_ask")
    @patch("app.match_person_in_prompt")
    def test_qa_person_ambiguity_still_stops_before_agent(self, match, ask):
        match.return_value = {"prompt": "List Sam's work", "status": "ambiguous", "ambiguous": [{"mention": "Sam"}]}
        response = self.client.post("/api/query", json={"question": "List Sam's work"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["needs_person_choice"])
        ask.assert_not_called()

    @patch("app.agent_ask")
    @patch("app.apply_person_choices")
    def test_qa_choice_roundtrip_retains_graph_chart_prediction_and_history(self, choices, ask):
        choices.return_value = {"prompt": "List Sample, Person's work", "ambiguous": [], "status": "resolved"}
        ask.return_value = {
            "answer": "Done", "rows": [{"DID": "D1", "hours": 12}], "cypher": "RETURN 12 AS hours",
            "graph": {"nodes": [{"id": "D1"}], "links": []},
            "visualization": {"chart_type": "bar", "x": "DID", "y": ["hours"]},
            "prediction": {"p50_hours": 12},
        }
        response = self.client.post("/api/query", json={
            "question": "List Sam's work", "person_match": {"prompt": "Sam"},
            "person_choices": {"Sam": "Sample, Person"}, "history": [{"role": "user", "content": "Before"}],
        })
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["graph"]["nodes"][0]["id"], "D1")
        self.assertEqual(payload["viz"], {"kind": "bar", "chart": {"x": "DID", "y": "hours"}})
        self.assertEqual(payload["prediction"]["p50_hours"], 12)
        self.assertEqual(ask.call_args.kwargs["history"], [{"role": "user", "content": "Before"}])

    @patch("app.close_driver")
    def test_shared_driver_is_closed_only_at_lifespan_shutdown(self, close):
        with TestClient(app):
            close.assert_not_called()
        close.assert_called_once()


class ArtifactResolverTests(unittest.TestCase):
    def test_published_similarity_bytes_are_never_mutated(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "similarity.joblib"
            source.write_bytes(b"original cache")
            with patch.object(prediction, "DEFAULT_SIMILARITY_CACHE_PATH", source), \
                 patch.object(prediction, "_prediction_artifact_cache_dir", return_value=root / "runtime"):
                (root / "runtime").mkdir()
                resolved = prediction._resolve_prediction_artifact(source, source, "UNSET_URL", "unused")
                self.assertNotEqual(resolved, source)
                resolved.write_bytes(b"runtime updates")
                self.assertEqual(source.read_bytes(), b"original cache")

    def test_missing_default_downloads_but_custom_missing_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            default = Path(directory) / "model.joblib"
            with patch.object(prediction, "_prediction_artifact_cache_dir", return_value=Path(directory)), \
                 patch.object(prediction, "_download_prediction_artifact", return_value=default) as download:
                prediction._resolve_prediction_artifact(default, default, "UNSET_URL", "pinned-url")
                download.assert_called_once_with("pinned-url", default)
                with self.assertRaises(FileNotFoundError):
                    prediction._resolve_prediction_artifact(Path(directory) / "custom.joblib", default, "UNSET_URL", "url")

    def test_artifacts_are_pinned_to_destination_repository(self):
        self.assertIn("LUM16/DID_QA_FastAPI/published-artifacts-26af6b5", prediction.GITHUB_ARTIFACT_BASE_URL)
        self.assertNotIn("/main/", prediction.GITHUB_ARTIFACT_BASE_URL)

    def test_planning_refresh_helpers_do_not_close_shared_driver(self):
        import du_team_recommendation as teams
        import tlf_person_allocation as people
        for module in [prediction, teams, people]:
            driver = Mock()
            driver.session.return_value.__enter__ = Mock(return_value=driver.session.return_value)
            driver.session.return_value.__exit__ = Mock(return_value=False)
            with patch.object(module, "load_env"), patch.object(module, "get_driver", return_value=driver):
                module._read_query("RETURN 1")
                driver.close.assert_not_called()


if __name__ == "__main__":
    unittest.main()
