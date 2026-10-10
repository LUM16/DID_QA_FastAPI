"""Tests for snapshot-only TLF person allocation."""

from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from tlf_person_allocation import (
    DEFAULT_TLF_PERSON_SNAPSHOT_PATH,
    _resolve_snapshot_path,
    _grouped_targets,
    allocation_excel_bytes,
    allocate_tlf_people,
    allocate_uploaded_tlfs,
    load_person_tlf_snapshot,
    refresh_person_tlf_snapshot,
    resolve_team_lead_name,
    team_lead_candidates,
)


def evidence(person: str, role: str, did: str, title: str = "AE Summary") -> dict:
    return {
        "person": person, "team_lead_name": "Team A", "did": did,
        "completion_date": "2026-09-01", "tlf_name": title, "tlf_type": "T",
        "tlf_source": "ADAE",
        "generation": person if role == "generation" else "",
        "qc": person if role == "qc" else "",
    }


class TlfPersonAllocationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.workspace = Path.cwd() / ".test-tlf-person-allocation"
        self.workspace.mkdir(exist_ok=True)

    def tearDown(self) -> None:
        for path in self.workspace.glob("*"):
            path.unlink()
        self.workspace.rmdir()

    def test_snapshot_roundtrip(self) -> None:
        path = self.workspace / "snapshot.joblib"
        with patch(
            "tlf_person_allocation._read_query",
            side_effect=[
                [{"team_lead_name": "Team A"}],
                [evidence("Gen", "generation", "D1")],
                [{"person": "Gen", "active_did_count": 2}],
            ],
        ):
            saved = refresh_person_tlf_snapshot(path)
        history, workload, details = load_person_tlf_snapshot(path)
        self.assertEqual(history[0]["person"], "Gen")
        self.assertEqual(workload[0]["active_did_count"], 2)
        self.assertEqual(saved["history_row_count"], details["history_row_count"])

    def test_default_snapshot_uses_lfs_resolver(self) -> None:
        resolved = Path("C:/temporary/tlf-snapshot.joblib")
        with patch("tlf_person_allocation.Path.exists", return_value=True), patch(
            "tlf_person_allocation._resolve_prediction_artifact", return_value=resolved
        ):
            self.assertEqual(_resolve_snapshot_path(DEFAULT_TLF_PERSON_SNAPSHOT_PATH), resolved)

    def test_team_resolver_accepts_only_actual_candidate(self) -> None:
        good = resolve_team_lead_name("Team A", ["Team A"])
        self.assertEqual(good["team_lead_name"], "Team A")
        for name in ["team a", "Team", "Invented Team", ""]:
            with self.subTest(name=name):
                with self.assertRaisesRegex(ValueError, "exact Team_Lead_Name"):
                    resolve_team_lead_name(name, ["Team A"])
        with self.assertRaisesRegex(ValueError, "no Team_Lead_Name"):
            resolve_team_lead_name("Team A", [])

    def test_candidates_are_unique_sorted_snapshot_names(self) -> None:
        self.assertEqual(
            team_lead_candidates([
                {"team_lead_name": "Team B"},
                {"team_lead_name": "Team A"},
                {"team_lead_name": "Team A"}, {"team_lead_name": ""},
            ]),
            ["Team A", "Team B"],
        )

    def test_uploaded_allocation_uses_no_llm_or_database_calls(self) -> None:
        rows = [evidence("Gen", "generation", "D1"), evidence("QC", "qc", "D2")]
        upload = BytesIO(
            b"Title,Type,Source Datasets\nAE Summary,T,ADAE\n"
        )
        upload.name = "scope.csv"
        data_upload = BytesIO(
            b"SDTM/ADaM,Domain/Dataset Name\nADaM,ADAE\n"
        )
        data_upload.name = "data.csv"
        with (
            patch("tlf_person_allocation.load_person_tlf_snapshot",
                  return_value=(rows, [], {"generated_at": "test"})),
            patch("tlf_person_allocation._resolve_recommendation_cache_path",
                  return_value=self.workspace / "upload-cache.joblib"),
            patch("vox_client.chat", side_effect=AssertionError("LLM called")) as llm,
            patch("tlf_person_allocation._read_query",
                  side_effect=AssertionError("Neo4j queried")) as query,
        ):
            result = allocate_uploaded_tlfs("Team A", upload, data_upload)
            with self.assertRaisesRegex(ValueError, "exact Team_Lead_Name"):
                allocate_uploaded_tlfs("Unknown", upload)
        self.assertEqual(result["team_match"], {"team_lead_name": "Team A"})
        self.assertEqual(result["allocations"][0]["generation"]["primary"]["person"], "Gen")
        self.assertEqual(result["allocations"][0]["qc"]["primary"]["person"], "QC")
        llm.assert_not_called()
        query.assert_not_called()

    def test_allocation_ui_requires_explicit_searchable_team_selection(self) -> None:
        from fastapi.testclient import TestClient
        from app import app

        rows = [
            evidence("Gen", "generation", "D1"),
            evidence("QC", "qc", "D2"),
            {**evidence("Other", "generation", "D3"), "team_lead_name": "Team B"},
        ]
        with (
            patch("tlf_person_allocation.load_person_tlf_snapshot",
                  return_value=(rows, [], {"generated_at": "test", "history_row_count": 3})),
            patch("tlf_person_allocation._resolve_recommendation_cache_path",
                  return_value=self.workspace / "ui-cache.joblib"),
            patch("planning_api.load_person_tlf_snapshot",
                  return_value=(rows, [], {"generated_at": "test", "history_row_count": 3})),
            patch("vox_client.chat", side_effect=AssertionError("LLM called")) as llm,
        ):
            client = TestClient(app)
            self.assertEqual(
                client.get("/api/planning/allocate/options").json()["team_leads"],
                ["Team A", "Team B"],
            )
            files = {
                "primary_file": ("scope.csv", b"Title,Type,Source Datasets\nAE Summary,T,ADAE\n", "text/csv"),
                "data_csv_file": ("data.csv", b"SDTM/ADaM,Domain/Dataset Name\nADaM,ADAE\n", "text/csv"),
            }
            invalid = client.post("/api/planning/allocate", data={"team_lead_name": "Partial"}, files=files)
            self.assertEqual(invalid.status_code, 422)
            response = client.post("/api/planning/allocate", data={"team_lead_name": "Team A"}, files=files)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["team_match"], {"team_lead_name": "Team A"})
            script = client.get("/static/planning.js").text
            self.assertIn("select.selectedIndex = mode === 'recommend' ? 0 : -1", script)
            self.assertIn("if (mode === 'allocate') { select.size = 5; label.append(search); }", script)
            self.assertIn("data.required = needsData; data.disabled = !needsData", script)
            self.assertNotIn("datalist", script)
            self.assertNotIn("DU", script)
        llm.assert_not_called()

    def test_distinct_primaries_and_two_backups(self) -> None:
        rows = [
            evidence("G1", "generation", "D1"), evidence("G2", "generation", "D2"),
            evidence("G3", "generation", "D3"), evidence("Q1", "qc", "D4"),
            evidence("Q2", "qc", "D5"), evidence("Q3", "qc", "D6"),
        ]
        result = allocate_tlf_people(
            [{"name": "AE Summary", "type": "T", "source": "ADAE"}],
            "Team A", rows, [], self.workspace / "cache.joblib",
        )
        allocation = result["allocations"][0]
        self.assertNotEqual(
            allocation["generation"]["primary"]["person"], allocation["qc"]["primary"]["person"]
        )
        self.assertEqual(len(allocation["generation"]["backups"]), 2)
        self.assertEqual(len(allocation["qc"]["backups"]), 2)

    def test_workload_and_balancing_change_later_primary(self) -> None:
        rows = [
            evidence("Busy", "generation", "D1", "AE Summary"),
            evidence("Free", "generation", "D2", "AE Summary"),
        ]
        workloads = [{"person": "Busy", "active_did_count": 10}, {"person": "Free", "active_did_count": 0}]
        result = allocate_tlf_people(
            [{"name": "AE Summary", "type": "T", "source": "ADAE"}],
            "Team A", rows, workloads, self.workspace / "cache.joblib",
        )
        self.assertEqual(result["allocations"][0]["generation"]["primary"]["person"], "Free")
        balanced = allocate_tlf_people(
            [
                {"name": "AE Summary", "type": "T", "source": "ADAE"},
                {"name": "Lab Summary", "type": "T", "source": "ADAE"},
            ],
            "Team A",
            [
                evidence("Alpha", "generation", "D3", "AE Summary"),
                evidence("Beta", "generation", "D4", "AE Summary"),
                evidence("Alpha", "generation", "D5", "Lab Summary"),
                evidence("Beta", "generation", "D6", "Lab Summary"),
            ],
            [],
            self.workspace / "balanced-cache.joblib",
        )
        self.assertEqual(balanced["allocations"][0]["generation"]["primary"]["person"], "Alpha")
        self.assertEqual(balanced["allocations"][1]["generation"]["primary"]["person"], "Beta")

    def test_same_person_qc_requires_review_when_no_distinct_qc_evidence(self) -> None:
        rows = [evidence("Only", "generation", "D1"), evidence("Only", "qc", "D2")]
        result = allocate_tlf_people(
            [{"name": "AE Summary", "type": "T", "source": "ADAE"}],
            "Team A", rows, [], self.workspace / "cache.joblib",
        )
        allocation = result["allocations"][0]
        self.assertTrue(allocation["lead_review_required"])
        self.assertEqual(allocation["generation"]["primary"]["person"], "Only")
        self.assertEqual(allocation["qc"]["primary"]["person"], "Only")

    def test_grouped_primaries_keep_roles_consistent(self) -> None:
        rows = [
            evidence("Gen", "generation", "D1", "AE Summary"),
            evidence("Gen", "generation", "D2", "AE Listing"),
            evidence("QC", "qc", "D3", "AE Summary"),
            evidence("QC", "qc", "D4", "AE Listing"),
            evidence("Other", "qc", "D5", "AE Summary"),
        ]
        result = allocate_tlf_people(
            [
                {"name": "AE Summary", "type": "T", "source": "AE"},
                {"name": "AE Listing", "type": "T", "source": "AE"},
            ],
            "Team A",
            rows,
            [],
            self.workspace / "cache.joblib",
            group_keys=["SDTM: AE", "SDTM: AE"],
        )
        allocations = result["allocations"]
        self.assertEqual(
            {row["generation"]["primary"]["person"] for row in allocations}, {"Gen"}
        )
        self.assertEqual(len({row["qc"]["primary"]["person"] for row in allocations}), 1)
        self.assertNotIn(
            allocations[0]["generation"]["primary"]["person"],
            {row["qc"]["primary"]["person"] for row in allocations},
        )

    def test_grouping_and_excel_export(self) -> None:
        grouped = _grouped_targets(
            {
                "tlfs": [
                    {"name": "AE Summary", "type": "T", "source": "AE"},
                    {"name": "AE Listing", "type": "L", "source": "ADAE"},
                    {"name": "No source", "type": "T", "source": ""},
                ],
                "sdtms": [{"name": "AE"}],
            }
        )
        self.assertEqual([group for _, group in grouped], [
            "SDTM: AE", "Source: ADAE", "Individual: TLF 3",
        ])
        result = allocate_tlf_people(
            [{"name": "AE Summary", "type": "T", "source": "AE"}],
            "Team A",
            [
                evidence("G1", "generation", "D1"),
                evidence("G2", "generation", "D2"),
                evidence("G3", "generation", "D3"),
                evidence("Q1", "qc", "D4"),
                evidence("Q2", "qc", "D5"),
                evidence("Q3", "qc", "D6"),
            ],
            [],
            self.workspace / "cache.xlsx.joblib",
            group_keys=["SDTM: AE"],
        )
        workbook = load_workbook(BytesIO(allocation_excel_bytes(result)))
        sheet = workbook["TLF Allocation"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.max_row, 2)
        self.assertIn("Group Type", headers)
        self.assertIn("Group Key", headers)
        self.assertIn("Generation Backup 2", headers)
        self.assertIn("QC Backup 2", headers)


if __name__ == "__main__":
    unittest.main()
