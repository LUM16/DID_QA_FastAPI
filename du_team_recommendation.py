"""Recommend DU teams for an uploaded delivery scope without training a model."""

from __future__ import annotations

import argparse
import csv
import io
import os
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import joblib

from effort_prediction import (
    DEFAULT_SIMILARITY_CACHE_PATH,
    DEFAULT_SIMILARITY_CACHE_URL,
    GITHUB_ARTIFACT_BASE_URL,
    TLF_SEMANTIC_MATCH_THRESHOLD,
    _cached_tlf_semantic_similarity,
    _load_similarity_cache,
    _normalized_name,
    _resolve_prediction_artifact,
    _save_similarity_cache,
    _similarity_candidates,
)
from neo4j_client import get_driver, load_env

TEAM_HISTORY_QUERY = """
MATCH (p:Person)-[:WORKS_ON]->(d:Delivery)
WHERE p.Team_Lead_Name IS NOT NULL
  AND trim(toString(p.Team_Lead_Name)) <> ''
  AND toLower(toString(d.DID_Status)) = 'completed'
  AND d.DID IS NOT NULL
  AND d.Actual_Delivery_Date IS NOT NULL
WITH DISTINCT toString(p.Team_Lead_Name) AS du_team, d
CALL (du_team) {
  MATCH (member:Person)
  WHERE toString(member.Team_Lead_Name) = du_team
    AND member.Group_Lead_Name IS NOT NULL
    AND trim(toString(member.Group_Lead_Name)) <> ''
  RETURN collect(DISTINCT toString(member.Group_Lead_Name)) AS group_leads
}
CALL (d) {
  OPTIONAL MATCH (d)-[t:HAS_TLF]->(tlf:TLF)
  RETURN collect(DISTINCT {
    name: tlf.Name, type: tlf.Type, source: tlf.Source,
    generation: t.Generation, qc: t.QC
  }) AS tlfs
}
RETURN du_team,
       group_leads,
       toString(d.DID) AS did,
       substring(toString(d.Actual_Delivery_Date), 0, 10) AS completion_date,
       tlfs
"""

TEAM_WORKLOAD_QUERY = """
MATCH (p:Person)-[:WORKS_ON]->(d:Delivery)
WHERE p.Team_Lead_Name IS NOT NULL
  AND trim(toString(p.Team_Lead_Name)) <> ''
  AND toLower(toString(d.DID_Status)) IN ['ongoing', 'planned']
RETURN p.Team_Lead_Name AS du_team,
       count(DISTINCT d) AS active_did_count,
       count(DISTINCT p) AS active_people_count
"""

SEMANTIC_DID_CANDIDATE_LIMIT = 50
DU_HISTORY_SNAPSHOT_VERSION = "du-team-history-v2-tlf-only-groups"
DEFAULT_DU_HISTORY_SNAPSHOT_PATH = (
    Path(__file__).resolve().parent / "artifacts" / "du_team_history_snapshot.joblib"
)
DEFAULT_DU_HISTORY_SNAPSHOT_URL = (
    f"{GITHUB_ARTIFACT_BASE_URL}/du_team_history_snapshot.joblib"
)


def _resolve_recommendation_cache_path(cache_path: Path) -> Path:
    """Resolve the default shared cache when Git-backed deployments receive an LFS pointer."""
    if cache_path.resolve() != DEFAULT_SIMILARITY_CACHE_PATH.resolve():
        return cache_path
    return _resolve_prediction_artifact(
        cache_path,
        DEFAULT_SIMILARITY_CACHE_PATH,
        "DID_EFFORT_SIMILARITY_CACHE_URL",
        DEFAULT_SIMILARITY_CACHE_URL,
    )


def _read_query(query: str) -> list[dict[str, Any]]:
    load_env()
    driver = get_driver()
    database = os.environ.get("NEO4J_DATABASE", "neo4j")
    with driver.session(database=database) as session:
        return session.execute_read(lambda tx: tx.run(query).data())


def _history_snapshot_path() -> Path:
    configured_path = os.environ.get("DU_TEAM_HISTORY_SNAPSHOT_PATH")
    return Path(configured_path) if configured_path else DEFAULT_DU_HISTORY_SNAPSHOT_PATH


def _resolve_history_snapshot_path(snapshot_path: Path) -> Path:
    """Resolve the default snapshot when Git-backed deployments receive an LFS pointer."""
    if snapshot_path.resolve() != DEFAULT_DU_HISTORY_SNAPSHOT_PATH.resolve():
        return snapshot_path
    return _resolve_prediction_artifact(
        snapshot_path,
        DEFAULT_DU_HISTORY_SNAPSHOT_PATH,
        "DU_TEAM_HISTORY_SNAPSHOT_URL",
        DEFAULT_DU_HISTORY_SNAPSHOT_URL,
    )


def refresh_history_snapshot(
    output_path: Path = DEFAULT_DU_HISTORY_SNAPSHOT_PATH,
) -> dict[str, Any]:
    """Fetch current Neo4j DU evidence and save it as a reusable local snapshot."""
    history_rows = _read_query(TEAM_HISTORY_QUERY)
    if not history_rows:
        raise ValueError("Neo4j returned no completed DU team history to cache.")
    workload_rows = _read_query(TEAM_WORKLOAD_QUERY)
    snapshot = {
        "version": DU_HISTORY_SNAPSHOT_VERSION,
        "generated_at": datetime.now().astimezone().isoformat(),
        "history_rows": history_rows,
        "workload_rows": workload_rows,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(output_path.suffix + ".tmp")
    joblib.dump(snapshot, temporary_path)
    temporary_path.replace(output_path)
    return {
        "path": str(output_path.resolve()),
        "generated_at": snapshot["generated_at"],
        "history_row_count": len(history_rows),
        "workload_row_count": len(workload_rows),
    }


def load_history_snapshot(
    snapshot_path: Path | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Load a manually refreshed DU history snapshot without querying Neo4j."""
    path = _resolve_history_snapshot_path(snapshot_path or _history_snapshot_path())
    if not path.exists():
        raise FileNotFoundError(
            f"DU history snapshot is unavailable: {path.resolve()}. "
            "Run `python du_team_recommendation.py refresh-history` after Neo4j updates."
        )
    snapshot = joblib.load(path)
    if (
        not isinstance(snapshot, dict)
        or snapshot.get("version") != DU_HISTORY_SNAPSHOT_VERSION
        or not isinstance(snapshot.get("history_rows"), list)
        or not isinstance(snapshot.get("workload_rows"), list)
        or not isinstance(snapshot.get("generated_at"), str)
    ):
        raise ValueError(f"Invalid DU history snapshot: {path.resolve()}")
    return (
        snapshot["history_rows"],
        snapshot["workload_rows"],
        {
            "path": str(path.resolve()),
            "generated_at": snapshot["generated_at"],
            "history_row_count": len(snapshot["history_rows"]),
            "workload_row_count": len(snapshot["workload_rows"]),
        },
    )


def _column(row: dict[str, Any], name: str) -> str:
    normalized = {
        "".join(character for character in key.lower() if character.isalnum()): value
        for key, value in row.items()
    }
    return str(normalized.get("".join(character for character in name.lower() if character.isalnum()), "") or "").strip()


def _split_sources(value: str) -> list[str]:
    return [item.strip() for item in value.replace(";", ",").split(",") if item.strip()]


def _scope_from_rows(
    tlf_rows: Iterable[dict[str, Any]], data_rows: Iterable[dict[str, Any]]
) -> dict[str, list[dict[str, str]]]:
    tlfs = []
    adams = []
    sdtms = []
    for row in tlf_rows:
        title = _column(row, "Title")
        if not title:
            continue
        source = _column(row, "Source Datasets")
        tlfs.append(
            {
                "name": title,
                "type": _column(row, "Type"),
                "source": source,
            }
        )
        for dataset in _split_sources(source):
            adams.append({"name": dataset})
    for row in data_rows:
        kind = _column(row, "SDTM/ADaM")
        name = _column(row, "Domain/Dataset Name")
        if not name:
            continue
        if _normalized_name(kind) == "ADAM":
            adams.append({"name": name})
        elif _normalized_name(kind) == "SDTM":
            sdtms.append({"name": name})
    return {
        "tlfs": _unique_items(tlfs),
        "adams": _unique_items(adams),
        "sdtms": _unique_items(sdtms),
    }


def group_lead_candidates(history_rows: Iterable[dict[str, Any]]) -> list[str]:
    """Return stored Group_Lead_Name values available for recommendation filtering."""
    return sorted(
        {
            str(group_lead).strip()
            for row in history_rows
            for group_lead in row.get("group_leads", [])
            if str(group_lead).strip()
        },
        key=str.casefold,
    )


def resolve_group_lead_name(
    entered_name: str, candidates: Iterable[str]
) -> str:
    """Resolve a uniquely identifiable Group Lead without inventing a name."""
    entered_tokens = set(re.findall(r"[a-z0-9]+", entered_name.casefold()))
    if not entered_tokens:
        raise ValueError("Enter a Group Lead name, for example Maggie.")
    matches = [
        candidate
        for candidate in candidates
        if entered_tokens
        <= set(re.findall(r"[a-z0-9]+", candidate.casefold()))
    ]
    if len(matches) != 1:
        raise ValueError(
            "The entered Group Lead name must uniquely match a Group_Lead_Name "
            "stored in the DU history snapshot."
        )
    return matches[0]


def _unique_items(items: Iterable[dict[str, Any]]) -> list[dict[str, str]]:
    unique: dict[tuple[str, str, str], dict[str, str]] = {}
    for item in items:
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        normalized = (
            _normalized_name(name),
            _normalized_name(item.get("type")),
            _normalized_name(item.get("source")),
        )
        unique[normalized] = {
            "name": name,
            "type": str(item.get("type") or "").strip(),
            "source": str(item.get("source") or "").strip(),
        }
    return list(unique.values())


def load_uploaded_scope(
    primary_file: Any,
    data_csv_file: Any | None = None,
    require_data: bool = True,
) -> dict[str, list[dict[str, str]]]:
    """Read TLF scope, requiring Data only for workflows that use it."""
    filename = str(getattr(primary_file, "name", "")).lower()
    content = primary_file.getvalue()
    if filename.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel upload support requires openpyxl.") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets = {sheet.title.strip().lower(): sheet for sheet in workbook.worksheets}
        if "tlf" not in sheets or (require_data and "data" not in sheets):
            required_sheets = "TLF and Data" if require_data else "TLF"
            raise ValueError(f"The Excel workbook must contain a sheet named {required_sheets}.")

        def rows(sheet_name: str) -> list[dict[str, Any]]:
            values = sheets[sheet_name].iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(values, ())]
            return [
                dict(zip(headers, values_row))
                for values_row in values
                if any(value is not None and str(value).strip() for value in values_row)
            ]

        return _scope_from_rows(rows("tlf"), rows("data") if "data" in sheets else [])
    if not filename.endswith(".csv") or (require_data and data_csv_file is None):
        required_files = "both TLF CSV and Data CSV files" if require_data else "a TLF CSV file"
        raise ValueError(f"Upload one .xlsx workbook, or {required_files}.")
    tlf_rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    data_rows = (
        list(csv.DictReader(io.StringIO(data_csv_file.getvalue().decode("utf-8-sig"))))
        if data_csv_file is not None
        else []
    )
    return _scope_from_rows(tlf_rows, data_rows)


def _item_names(items: Iterable[dict[str, Any]]) -> set[str]:
    return {_normalized_name(item.get("name")) for item in items if item.get("name")}


def _date_or_minimum(value: Any) -> date:
    try:
        return datetime.fromisoformat(str(value)[:10]).date()
    except ValueError:
        return date.min


def _recency_score(completion_date: Any) -> float:
    days = (date.today() - _date_or_minimum(completion_date)).days
    if days <= 180:
        return 1.0
    if days <= 365:
        return 0.7
    if days <= 730:
        return 0.4
    return 0.15 if completion_date else 0.0


def _tlf_union_coverage(
    target_tlfs: list[dict[str, str]],
    history: list[dict[str, Any]],
    semantic_candidates: Iterable[dict[str, Any]],
    cache: dict[str, Any],
) -> tuple[float, list[dict[str, Any]]]:
    """Match each target TLF to its best available historical DID evidence."""
    candidate_by_did = {
        str(record["did"]): record for record in semantic_candidates
    }
    evidence = []
    for target_tlf in target_tlfs:
        target_title = _normalized_name(target_tlf.get("name"))
        target_words = {
            word for word in target_title.split() if len(word) >= 3
        }
        for record in history:
            if any(
                _normalized_name(item.get("name")) == target_title
                or target_words
                & {
                    word
                    for word in _normalized_name(item.get("name")).split()
                    if len(word) >= 3
                }
                for item in record["tlfs"]
            ):
                candidate_by_did[str(record["did"])] = record
        best: dict[str, Any] | None = None
        for record in candidate_by_did.values():
            similarity, single_tlf_coverage = _cached_tlf_semantic_similarity(
                [target_tlf], record["tlfs"], cache
            )
            candidate = {
                "did": record["did"],
                "completion_date": record.get("completion_date"),
                "matched": single_tlf_coverage == 1.0,
                "similarity": similarity,
            }
            if best is None or (
                candidate["matched"],
                candidate["similarity"],
                _date_or_minimum(candidate["completion_date"]),
            ) > (
                best["matched"],
                best["similarity"],
                _date_or_minimum(best["completion_date"]),
            ):
                best = candidate
        evidence.append(
            {
                "target_tlf": target_tlf["name"],
                "matched": bool(best and best["matched"]),
                "evidence_did": best["did"] if best and best["matched"] else None,
                "completion_date": (
                    best["completion_date"] if best and best["matched"] else None
                ),
            }
        )
    coverage = (
        sum(item["matched"] for item in evidence) / len(target_tlfs)
        if target_tlfs
        else 1.0
    )
    return coverage, evidence


def recommend_teams(
    scope: dict[str, list[dict[str, str]]],
    history_rows: Iterable[dict[str, Any]],
    workload_rows: Iterable[dict[str, Any]],
    cache_path: Path = DEFAULT_SIMILARITY_CACHE_PATH,
    top_n: int = 3,
    group_lead_name: str | None = None,
) -> dict[str, Any]:
    """Rank current DU teams from completed Delivery scope and active workload."""
    resolved_cache_path = _resolve_recommendation_cache_path(cache_path)
    target = {"tlfs": _unique_items(scope.get("tlfs", [])), "adams": [], "sdtms": []}
    if not target["tlfs"]:
        raise ValueError("No TLF scope was found in the uploaded file.")
    by_team: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in history_rows:
        team = str(row.get("du_team") or "").strip()
        did = str(row.get("did") or "").strip()
        group_leads = [
            str(item).strip()
            for item in row.get("group_leads", [])
            if str(item).strip()
        ]
        if (
            not team
            or not did
            or (group_lead_name is not None and group_lead_name not in group_leads)
        ):
            continue
        by_team[team][did] = {
            **row,
            "group_leads": group_leads,
            "tlfs": _unique_items(row.get("tlfs") or []),
        }
    if not by_team:
        raise ValueError("Neo4j returned no completed DU team history.")

    workload = {
        str(row["du_team"]): int(row.get("active_did_count") or 0)
        for row in workload_rows
        if row.get("du_team")
    }
    maximum_workload = max(workload.values(), default=0)
    cache = _load_similarity_cache(resolved_cache_path)
    recommendations = []
    for team, histories in by_team.items():
        history = list(histories.values())
        historical_similarity = []
        candidate_history = sorted(
            _similarity_candidates(target, history),
            key=lambda record: _date_or_minimum(record.get("completion_date")),
            reverse=True,
        )[:SEMANTIC_DID_CANDIDATE_LIMIT]
        tlf_coverage, tlf_coverage_evidence = _tlf_union_coverage(
            target["tlfs"], history, candidate_history, cache
        )
        for record in candidate_history:
            tlf_similarity, single_did_tlf_coverage = _cached_tlf_semantic_similarity(
                target["tlfs"], record["tlfs"], cache
            )
            combined = single_did_tlf_coverage
            historical_similarity.append(
                {
                    "did": record["did"],
                    "completion_date": record.get("completion_date"),
                    "combined_similarity": combined,
                    "tlf_semantic_coverage": single_did_tlf_coverage,
                    "tlf_semantic_similarity": tlf_similarity,
                }
            )
        historical_similarity.sort(
            key=lambda row: (row["combined_similarity"], _date_or_minimum(row["completion_date"])),
            reverse=True,
        )
        best = historical_similarity[0] if historical_similarity else None
        similar_70 = [
            row for row in historical_similarity
            if row["combined_similarity"] >= TLF_SEMANTIC_MATCH_THRESHOLD
        ]
        similar_85 = [
            row for row in historical_similarity if row["combined_similarity"] >= 0.85
        ]
        similar_score = (
            0.6 * (best["combined_similarity"] if best else 0.0)
            + 0.4 * min(1.0, len(similar_70) / 5)
        )
        recent_score = max(
            (_recency_score(row["completion_date"]) for row in similar_70),
            default=0.0,
        )
        active_dids = workload.get(team, 0)
        workload_score = (
            1.0 - active_dids / maximum_workload if maximum_workload else 1.0
        )
        score = (
            37 * tlf_coverage
            + 18 * similar_score
            + 10 * recent_score
            + 35 * workload_score
        )
        recommendations.append(
            {
                "du_team": team,
                "score": round(score, 1),
                "tlf_semantic_coverage": round(tlf_coverage, 3),
                "tlf_coverage_evidence": tlf_coverage_evidence,
                "similar_did_count_ge_70": len(similar_70),
                "similar_did_count_ge_85": len(similar_85),
                "recent_experience_score": round(recent_score, 3),
                "active_did_count": active_dids,
                "completed_did_count": len(history),
                "semantic_candidate_did_count": len(candidate_history),
                "similar_dids": historical_similarity[:5],
            }
        )
    recommendations.sort(key=lambda row: (-row["score"], row["du_team"]))
    for rank, row in enumerate(recommendations, start=1):
        row["rank"] = rank
        row["recommendation_level"] = (
            "Recommended" if row["score"] >= 80
            else "Suitable with review" if row["score"] >= 65
            else "Backup option" if row["score"] >= 45
            else "Insufficient evidence"
        )
    _save_similarity_cache(cache, resolved_cache_path)
    return {
        "target_counts": {"tlfs": len(target["tlfs"])},
        "group_lead_name": group_lead_name,
        "recommendations": recommendations[:top_n],
        "cache": {
            "hits": cache["hits"],
            "misses": cache["misses"],
            "candidate_pairs": cache["hits"] + cache["misses"],
        },
    }


def recommend_uploaded_scope(
    group_lead_input: str, primary_file: Any, data_csv_file: Any | None = None
) -> dict[str, Any]:
    """Load an uploaded scope and rank it against the locally cached DU history."""
    history_rows, workload_rows, snapshot = load_history_snapshot()
    group_lead_name = resolve_group_lead_name(
        group_lead_input, group_lead_candidates(history_rows)
    )
    result = recommend_teams(
        load_uploaded_scope(primary_file, data_csv_file, require_data=False),
        history_rows,
        workload_rows,
        group_lead_name=group_lead_name,
    )
    result["history_snapshot"] = snapshot
    return result


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Refresh the local Neo4j evidence snapshot used by DU recommendations."
    )
    subcommands = parser.add_subparsers(dest="command", required=True)
    refresh_parser = subcommands.add_parser(
        "refresh-history",
        help="Query Neo4j once and write the DU history snapshot used by the UI.",
    )
    refresh_parser.add_argument(
        "--output",
        type=Path,
        default=_history_snapshot_path(),
        help="Snapshot output path (default: artifacts/du_team_history_snapshot.joblib).",
    )
    arguments = parser.parse_args()
    if arguments.command == "refresh-history":
        details = refresh_history_snapshot(arguments.output)
        print(
            "DU history snapshot refreshed: "
            f"{details['history_row_count']:,} history rows, "
            f"{details['workload_row_count']:,} workload rows."
        )
        print(f"Generated at: {details['generated_at']}")
        print(f"Path: {details['path']}")


if __name__ == "__main__":
    main()
